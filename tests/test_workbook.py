"""Package checks and VBA tests in pyOpenVBA's interpreter (not native Excel)."""

from __future__ import annotations

import io
from pathlib import Path
import re
import struct
import sys
import unittest
from unittest.mock import patch
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / ".build" / "deps"))

import olefile
import openpyxl
from pyopenvba import ExcelFile
from pyopenvba.apps.excel import ExcelApplication
from pyopenvba.exceptions import VBARuntimeError
from pyopenvba.interpreter import EMPTY, parse_module

from tools.build_workbook import COLORS, SHEETS, STATUSES, build

OUT = ROOT / ".build" / "tests"
EMPTY_BOOK = OUT / "test_empty.xlsm"
SAMPLE_BOOK = OUT / "test_sample.xlsm"


def decompress_source(data: bytes) -> bytes:
    """Independent MS-OVBA decompressor used only to verify embedded source."""
    if data[0] != 1:
        raise ValueError("Invalid compressed container")
    position = 1
    result = bytearray()
    while position < len(data):
        header = struct.unpack_from("<H", data, position)[0]
        end = position + (header & 0xFFF) + 3
        position += 2
        if ((header >> 12) & 7) != 3:
            raise ValueError("Invalid chunk header")
        if not header & 0x8000:
            result.extend(data[position:end])
            position = end
            continue
        chunk_start = len(result)
        while position < end:
            flags = data[position]
            position += 1
            for bit in range(8):
                if position >= end:
                    break
                if flags & (1 << bit):
                    token = struct.unpack_from("<H", data, position)[0]
                    position += 2
                    offset_bits = max(4, (len(result) - chunk_start - 1).bit_length())
                    length_bits = 16 - offset_bits
                    length = (token & ((1 << length_bits) - 1)) + 3
                    offset = (token >> length_bits) + 1
                    for _ in range(length):
                        result.append(result[-offset])
                else:
                    result.append(data[position])
                    position += 1
    return bytes(result)


def setUpModule():
    build(EMPTY_BOOK)
    build(SAMPLE_BOOK, sample=True)


class PackageTests(unittest.TestCase):
    def test_xml_integrity_and_macro_content_type(self):
        for path in (EMPTY_BOOK, SAMPLE_BOOK):
            with self.subTest(path=path.name), ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
                for name in archive.namelist():
                    if name.endswith((".xml", ".rels")):
                        ET.fromstring(archive.read(name))
                content_types = archive.read("[Content_Types].xml")
                self.assertIn(b"application/vnd.ms-excel.sheet.macroEnabled.main+xml", content_types)
                self.assertIn("xl/vbaProject.bin", archive.namelist())
                self.assertFalse(any("externalLink" in name for name in archive.namelist()))

    def test_sheets_protection_font_and_validation(self):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            book = openpyxl.load_workbook(EMPTY_BOOK, keep_vba=True)
        try:
            self.assertEqual(book.sheetnames, [name for name, _ in SHEETS])
            for name, code in SHEETS:
                self.assertEqual(book[name].sheet_properties.codeName, code)
                self.assertTrue(book[name].protection.sheet)
            self.assertEqual(book["設定"].sheet_state, "veryHidden")
            for font in book._fonts:
                self.assertEqual(font.name, "Meiryo UI")
            editor = book["付箋の編集"]
            for address in ("B6", "B9", "B14", "B24"):
                self.assertFalse(editor[address].protection.locked)
                self.assertEqual(editor[address].number_format, "@")
            self.assertTrue(editor["B4"].protection.locked)
            validation = next(v for v in editor.data_validations.dataValidation if v.type == "list")
            self.assertEqual(validation.formula1, "StatusList")
            self.assertEqual([book["設定"].cell(row, 1).value for row in range(1, 6)], list(STATUSES))
            self.assertIsNone(book["タスク一覧"]["A7"].value)
        finally:
            book.close()
            book.vba_archive.close()

    def test_all_buttons_resolve_to_public_macros(self):
        ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        with ExcelFile(SAMPLE_BOOK) as host:
            source = host.get_module("TaskBoard")
            macros = set(re.findall(r"Public Sub (\w+)\(", source))
        found = []
        with ZipFile(SAMPLE_BOOK) as archive:
            for name in archive.namelist():
                if re.fullmatch(r"xl/drawings/drawing\d+\.xml", name):
                    for shape in ET.fromstring(archive.read(name)).findall(".//xdr:sp", ns):
                        action = shape.get("macro", "").removeprefix("[0]!")
                        self.assertIn(action, macros)
                        found.append(action)
                        names = [node.get("typeface") for node in shape.findall(".//a:latin", ns)]
                        self.assertTrue(names)
                        self.assertEqual(set(names), {"Meiryo UI"})
        self.assertEqual(found.count("EditCard"), 5)
        self.assertIn("SaveTask", found)

    def test_vba_roundtrip_and_syntax(self):
        with ExcelFile(EMPTY_BOOK) as host, ZipFile(EMPTY_BOOK) as archive:
            project = host.vba_project()
            self.assertEqual(project.code_page, 932)
            self.assertEqual(set(project.module_names()), {code for _, code in SHEETS} | {"ThisWorkbook", "TaskRules", "TaskBoard"})
            with olefile.OleFileIO(io.BytesIO(archive.read("xl/vbaProject.bin"))) as ole:
                for module in project.modules:
                    raw = ole.openstream(["VBA", module.stream_name]).read()
                    decoded = decompress_source(raw[module.text_offset:]).decode("cp932")
                    self.assertEqual(decoded, module.source)
                    parse_module(decoded)
                    if module.name in ("TaskRules", "TaskBoard"):
                        original = (ROOT / "src" / f"{module.name}.bas").read_text(encoding="utf-8")
                        self.assertEqual(decoded.replace("\r\n", "\n"), original)


class VbaBehaviorTests(unittest.TestCase):
    def setUp(self):
        self.app = ExcelApplication.open(EMPTY_BOOK)
        self.read_only = False
        # The interpreter has no filesystem lock state. Mock only this host
        # property; execute the production VBA and its read-only guard intact.
        host_type = type(self.app.workbooks()[0])
        original_get = host_type.vba_get

        def host_get(host, name, args=None, named=None):
            if name.casefold() == "readonly":
                return self.read_only
            return original_get(host, name, args, named)

        mocked = patch.object(host_type, "vba_get", host_get)
        mocked.start()
        self.addCleanup(mocked.stop)
        self.app.run("ProtectSheets")

    def cell(self, row, col):
        return self.app.evaluate(f"TaskData.Cells({row}, {col}).Value2")

    def store(self, status="起票済み", task_id="新規", title="日本語のタスク", body="本文\n２行目", notes="関連実装：src/sample.py"):
        return self.app.run("StoreTask", task_id, title, body, notes, status)

    def test_status_colors_and_unknown_status(self):
        for index, (status, color) in enumerate(zip(STATUSES, COLORS)):
            self.assertEqual(self.app.run("StatusIndex", status), index)
            rgb = bytes.fromhex(color[1:])
            self.assertEqual(self.app.run("StatusColor", status), rgb[0] + 256 * rgb[1] + 65536 * rgb[2])
        self.assertEqual(self.app.run("StatusIndex", "未定義"), -1)

    def test_validation_before_mutation(self):
        for title, body, notes, status in ((" \t\n　", "", "", "起票済み"), ("a" * 121, "", "", "対応中"), ("題", "a" * 8001, "", "対応中"), ("題", "", "a" * 8001, "対応中"), ("題", "", "", "unknown")):
            with self.subTest(title=title[:10], status=status):
                with self.assertRaises(VBARuntimeError):
                    self.store(status=status, title=title, body=body, notes=notes)
                self.assertEqual(self.app.run("LastTaskRow"), 6)
        self.store(title="題" * 120, body="内" * 8000, notes="備" * 8000)
        self.assertEqual(self.app.run("LastTaskRow"), 7)

    def test_create_update_and_japanese_multiline_preserved(self):
        task_id = self.store()
        self.assertEqual(task_id, "T000001")
        self.assertEqual(self.cell(7, 3), "本文\n２行目")
        self.assertEqual(self.cell(7, 4), "関連実装：src/sample.py")
        created = self.cell(7, 6)
        for status in STATUSES:
            self.assertEqual(self.store(status=status, task_id=task_id), task_id)
            self.assertEqual(self.app.run("LastTaskRow"), 7)
            self.assertEqual(self.cell(7, 5), status)
            self.assertEqual(self.cell(7, 6), created)
        self.assertTrue(self.cell(7, 8))

    def test_completion_timestamp_preserved_then_cleared_on_reopen(self):
        task_id = self.store(status="完了")
        completed = self.cell(7, 8)
        self.store(status="完了", task_id=task_id, notes="完了後に追記")
        self.assertEqual(self.cell(7, 8), completed)
        self.store(status="対応中", task_id=task_id)
        self.assertIn(self.cell(7, 8), (None, "", 0))
        self.assertEqual(self.cell(7, 5), "対応中")

    def test_ids_and_missing_update_do_not_overwrite_another_task(self):
        self.assertEqual(self.store(), "T000001")
        self.assertEqual(self.store(), "T000002")
        with self.assertRaises(VBARuntimeError):
            self.store(task_id="T999999")
        self.assertEqual(self.app.run("LastTaskRow"), 8)
        self.assertEqual(self.cell(7, 1), "T000001")
        self.assertEqual(self.cell(8, 1), "T000002")

    def test_read_only_workbook_does_not_change_records(self):
        self.read_only = True
        with self.assertRaises(VBARuntimeError):
            self.store()
        self.assertEqual(self.app.run("LastTaskRow"), 6)

    def test_formula_like_input_is_stored_as_text(self):
        self.store(title="=1+1", body="+SUM(A1:A2)", notes="@test")
        for col, expected in ((2, "=1+1"), (3, "+SUM(A1:A2)"), (4, "@test")):
            self.assertEqual(self.cell(7, col), expected)
            self.assertFalse(self.app.evaluate(f"TaskData.Cells(7, {col}).HasFormula"))

    def test_board_archive_routing_and_reopen_with_mock_canvas(self):
        # pyOpenVBA does not implement native TextFrame2 formatting. Replace
        # only the drawing boundary and observe the destination/lane/slot that
        # the production RenderBoards routine chooses. Native rendering has a
        # separate Excel smoke test and is not claimed as verified here.
        drawn = []
        original_call = self.app.interpreter.call

        def record_draw(proc, module, args, named, **kwargs):
            if proc.name.casefold() == "addcard":
                drawn.append((args[0].name, int(args[1]), int(args[2]), int(args[3])))
                return EMPTY
            return original_call(proc, module, args, named, **kwargs)

        mocked = patch.object(self.app.interpreter, "call", record_draw)
        mocked.start()
        self.addCleanup(mocked.stop)
        task_id = self.store(status="対応中")
        self.app.run("RenderBoards")
        self.assertEqual(drawn, [("タスクボード", 7, 1, 0)])
        drawn.clear()
        self.store(status="完了", task_id=task_id)
        self.app.run("RenderBoards")
        self.assertEqual(drawn, [("完了", 7, 0, 0)])
        drawn.clear()
        self.store(status="起票済み", task_id=task_id)
        self.app.run("RenderBoards")
        self.assertEqual(drawn, [("タスクボード", 7, 0, 0)])


if __name__ == "__main__":
    unittest.main(verbosity=2)
